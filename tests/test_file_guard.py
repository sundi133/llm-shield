"""POST /guardrails/file — file-attachment screening through the input pipeline.

Covers the spec's edge cases: text/DOCX/XLSX/PDF extraction feeding the same
verdict pipeline as /guardrails/input, filename-only screening for unsupported
types, the 10MB cap (413), empty files, corrupt files and missing extraction
libs (fail-open with note, never 500), extraction truncation, and the audit /
metrics side effects incl. device_id.

    PYTHONPATH=. pytest tests/test_file_guard.py
"""

import fnmatch
import io
import json

import pytest
from unittest.mock import patch

from config.schema import ShieldConfig, GuardrailConfig


# Blocklist keyword used across fixtures; any file whose extracted text (or
# filename) contains it must produce a block verdict.
BAD = "default_bad"


@pytest.fixture
def classify_config():
    return ShieldConfig(
        guardrails={
            "keyword_blocklist": GuardrailConfig(
                enabled=True,
                action="block",
                settings={"keywords": [BAD], "case_insensitive": True},
            ),
        },
    )


@pytest.fixture
def app(classify_config):
    import config.schema as cs
    from guardrails import registry as reg

    original = cs.config
    cs.config = classify_config
    reg._registry.clear()
    reg._discovered = False
    with patch("config.schema.load_config", return_value=classify_config):
        from core.app import create_app
        app = create_app()
    yield app
    cs.config = original
    reg._registry.clear()
    reg._discovered = False


@pytest.fixture
def client(app):
    from starlette.testclient import TestClient
    return TestClient(app)


def _upload(client, filename, data, content_type="application/octet-stream",
            headers=None, form=None):
    return client.post(
        "/guardrails/file",
        files={"file": (filename, io.BytesIO(data), content_type)},
        data=form or {},
        headers=headers or {},
    )


# ── fixtures for binary formats ───────────────────────────────────────────

def _docx_bytes(text):
    import docx
    buf = io.BytesIO()
    d = docx.Document()
    d.add_paragraph(text)
    d.save(buf)
    return buf.getvalue()


def _xlsx_bytes(text):
    import openpyxl
    buf = io.BytesIO()
    wb = openpyxl.Workbook()
    wb.active["A1"] = text
    wb.save(buf)
    return buf.getvalue()


def _pdf_bytes(text):
    """Minimal valid single-page PDF with one text object (real xref/EOF)."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >> stream\n"
        + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objs, start=1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj ".encode() + obj + b" endobj\n")
    xref_pos = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n".encode())
    out.write(b"0000000000 65535 f \n")
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer << /Size {len(objs) + 1} /Root 1 0 R >>\n"
              f"startxref\n{xref_pos}\n%%EOF\n".encode())
    return out.getvalue()


# ── verdicts per file type ────────────────────────────────────────────────

def test_txt_block_and_pass(client):
    r = _upload(client, "notes.txt", f"contains {BAD} content".encode(), "text/plain")
    assert r.status_code == 200
    body = r.json()
    assert body["safe"] is False and body["action"] == "block"
    assert body["file"]["name"] == "notes.txt"
    assert body["file"]["extracted_chars"] > 0
    assert body["file"]["note"] is None

    r2 = _upload(client, "notes.txt", b"perfectly fine text", "text/plain")
    assert r2.json()["action"] == "pass"


def test_docx_extraction_feeds_pipeline(client):
    r = _upload(client, "report.docx", _docx_bytes(f"quarterly {BAD} numbers"))
    body = r.json()
    assert body["action"] == "block"
    assert body["file"]["extracted_chars"] > 0


def test_xlsx_extraction_feeds_pipeline(client):
    r = _upload(client, "sheet.xlsx", _xlsx_bytes(f"{BAD} cell value"))
    body = r.json()
    assert body["action"] == "block"
    assert body["file"]["extracted_chars"] > 0


def test_pdf_extraction_feeds_pipeline(client):
    r = _upload(client, "doc.pdf", _pdf_bytes(f"secret {BAD} data"), "application/pdf")
    body = r.json()
    assert body["action"] == "block"
    assert body["file"]["extracted_chars"] > 0


def test_unsupported_type_screens_filename_only(client):
    # keyword in the FILENAME of a non-extractable type must still trip policy
    r = _upload(client, f"{BAD}_customers.zip", b"\x50\x4b\x03\x04binary")
    body = r.json()
    assert body["action"] == "block"
    assert body["file"]["note"] == "content not screenable (unsupported type)"
    assert body["file"]["extracted_chars"] == 0

    # clean filename + unsupported type -> pass with the same note
    r2 = _upload(client, "photo.zip", b"\x50\x4b\x03\x04binary")
    assert r2.json()["action"] == "pass"
    assert r2.json()["file"]["note"] == "content not screenable (unsupported type)"


# ── limits and degraded paths ─────────────────────────────────────────────

def test_oversize_returns_413(client, monkeypatch):
    monkeypatch.setenv("SHIELD_FILE_MAX_BYTES", "100")
    r = _upload(client, "big.txt", b"x" * 200, "text/plain")
    assert r.status_code == 413


def test_missing_file_part_rejected(client):
    r = client.post("/guardrails/file", data={"session_id": "s1"})
    assert r.status_code == 422  # FastAPI validation: file is required


def test_empty_file_passes_with_note(client):
    r = _upload(client, "empty.txt", b"", "text/plain")
    body = r.json()
    assert r.status_code == 200
    assert body["action"] == "pass"
    assert body["file"]["note"] == "no extractable text"


def test_corrupt_docx_fails_open_with_note(client):
    r = _upload(client, "broken.docx", b"not a real docx at all")
    body = r.json()
    assert r.status_code == 200
    assert body["action"] == "pass"
    assert body["file"]["note"].startswith("extraction failed")


def test_missing_extraction_lib_fails_open(client, monkeypatch):
    import api.routes_classify as rc

    def boom(data, max_chars):
        raise ImportError("pypdf not installed")
    monkeypatch.setattr(rc, "_extract_pdf", boom)
    r = _upload(client, "doc.pdf", _pdf_bytes("whatever"), "application/pdf")
    body = r.json()
    assert r.status_code == 200
    assert body["action"] == "pass"
    assert body["file"]["note"] == "extraction unavailable (missing library)"


def test_extraction_truncation_flag(client, monkeypatch):
    monkeypatch.setenv("SHIELD_FILE_EXTRACT_MAX_CHARS", "10")
    r = _upload(client, "long.txt", b"a" * 100, "text/plain")
    body = r.json()
    assert body["file"]["extraction_truncated"] is True
    assert body["file"]["extracted_chars"] == 10


# ── files block only on real policy violations ────────────────────────────

def test_warn_only_guardrail_demoted_on_files(client, monkeypatch):
    """A block from a warn-only guardrail (default: adversarial_detection;
    keyword used here as the stand-in since it blocks deterministically)
    becomes a warn on /guardrails/file — file passes with visibility."""
    monkeypatch.setenv("SHIELD_FILE_WARN_ONLY_GUARDRAILS", "keyword_blocklist")
    r = _upload(client, "paper.txt", f"discusses {BAD} responsibly".encode(), "text/plain")
    body = r.json()
    assert body["safe"] is True
    assert body["action"] == "warn"
    kw = next(g for g in body["guardrail_results"] if g["guardrail"] == "keyword_blocklist")
    assert kw["action"] == "warn"
    assert "demoted to warn for file screening" in kw["message"]


def test_warn_only_escape_hatch_restores_blocking(client, monkeypatch):
    monkeypatch.setenv("SHIELD_FILE_WARN_ONLY_GUARDRAILS", "")
    r = _upload(client, "notes.txt", f"contains {BAD}".encode(), "text/plain")
    assert r.json()["action"] == "block"


def test_demotion_does_not_affect_other_guardrails(client, monkeypatch):
    """Default warn-only set is adversarial_detection only — keyword blocks
    still block files (this is the existing behavior test, re-asserted with
    the demotion code in place)."""
    monkeypatch.delenv("SHIELD_FILE_WARN_ONLY_GUARDRAILS", raising=False)
    r = _upload(client, "notes.txt", f"contains {BAD}".encode(), "text/plain")
    assert r.json()["action"] == "block"


def test_demote_recomputes_root_action_with_mixed_failures():
    """Demoted guardrail + a real block -> root action stays block."""
    from api.routes_classify import _demote_file_verdicts
    result = {
        "safe": False, "action": "block",
        "guardrail_results": [
            {"guardrail": "adversarial_detection", "passed": False, "action": "block", "message": "Unsafe [none]"},
            {"guardrail": "pii_detection", "passed": False, "action": "block", "message": "SSN found"},
        ],
    }
    out = _demote_file_verdicts(result)
    assert out["action"] == "block" and out["safe"] is False
    adv = next(g for g in out["guardrail_results"] if g["guardrail"] == "adversarial_detection")
    assert adv["action"] == "warn"


def test_adversarial_untyped_verdict_warns_not_blocks(monkeypatch):
    """is_adversarial=true with attack_type=none is a contradiction — warn.
    A typed attack still blocks; the escape hatch restores old behavior."""
    import asyncio
    import guardrails.input.adversarial as adv_mod

    async def fake_llm(csv):
        async def call(**kwargs):
            return {"choices": [{"message": {"content": csv}}]}
        return call

    guard = adv_mod.AdversarialGuardrail.__new__(adv_mod.AdversarialGuardrail)
    guard.name = "adversarial_detection"
    guard._temp_config = {"action": "block"}  # configured_action reads this

    def run(csv):
        async def go():
            monkeypatch.setattr(adv_mod, "async_llm_call", await fake_llm(csv))
            return await guard._check_single("some text", [], 0.5)
        return asyncio.run(go())

    untyped = run("true,none,1.0")
    assert untyped.passed is False and untyped.action == "warn"
    assert "demoted to warn" in untyped.message

    typed = run("true,prompt_injection,0.9")
    assert typed.passed is False and typed.action == "block"

    monkeypatch.setenv("SHIELD_ADVERSARIAL_BLOCK_ON_UNTYPED", "true")
    untyped_hatch = run("true,none,1.0")
    assert untyped_hatch.action == "block"

    monkeypatch.delenv("SHIELD_ADVERSARIAL_BLOCK_ON_UNTYPED")
    clean = run("false,none,0.1")
    assert clean.passed is True


# ── guard-path integration invariants ─────────────────────────────────────

def test_file_endpoint_is_middleware_guarded():
    """Regression guard: /guardrails/file MUST be in ShieldMiddleware's
    guarded set or tenant_config/agent_key/tenant_id are never populated —
    tenant guardrails would silently not run on files and monitor-mode
    tenants would get enforce."""
    from core.middleware import ShieldMiddleware
    assert "/guardrails/file" in ShieldMiddleware._GUARDED_EXACT


def test_docx_content_type_without_extension(client):
    """Browsers set the OOXML content type reliably; extension may be absent."""
    r = client.post(
        "/guardrails/file",
        files={"file": ("upload", io.BytesIO(_docx_bytes(f"{BAD} inside")),
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert r.json()["action"] == "block"


# ── side effects: audit metadata + guardrail metrics ──────────────────────

class FakeRedis:
    def __init__(self):
        self.store = {}
        self.zsets = {}

    def hincrby(self, key, field, n=1):
        h = self.store.setdefault(key, {})
        h[field] = int(h.get(field, 0)) + n
        return h[field]

    def hincrbyfloat(self, key, field, n):
        h = self.store.setdefault(key, {})
        h[field] = float(h.get(field, 0)) + n
        return h[field]

    def expire(self, key, ttl):
        return True

    def hgetall(self, key):
        return {k: str(v) for k, v in self.store.get(key, {}).items()}

    def scan(self, cursor=0, match="*", count=100):
        return 0, [k for k in self.store if fnmatch.fnmatch(k, match)]

    def pipeline(self, *args, **kwargs):
        raise TypeError("no pipeline")

    def zadd(self, key, mapping):
        self.zsets.setdefault(key, []).extend(mapping.items())
        return len(mapping)


def test_audit_and_metrics_side_effects(client, monkeypatch):
    from storage import tenant_store
    import storage.audit_log as audit_log_mod

    r = FakeRedis()
    monkeypatch.setattr(tenant_store, "_get_redis", lambda: r)
    monkeypatch.setattr(tenant_store, "resolve_request_tenant_id",
                        lambda req: "bankco", raising=False)

    async def sync_log(self, entry):
        self._write_sync(entry)
    monkeypatch.setattr(audit_log_mod.AuditLogger, "log", sync_log)

    resp = _upload(client, "leak.txt", f"{BAD} data".encode(), "text/plain",
                   headers={"X-Device-Id": "LAPTOP-42",
                            "X-Agent-Key": "alice@co.com"},
                   form={"session_id": "sess-1"})
    assert resp.status_code == 200

    # guardrail metrics recorded for the tenant
    metric_keys = [k for k in r.store if k.startswith("guardrail:metrics:bankco:")]
    assert metric_keys, "metrics hash should be written"

    # audit entry with file metadata + device id
    entries = r.zsets.get("audit:bankco", [])
    assert entries, "audit entry should be written"
    record = json.loads(entries[-1][0])
    assert record["endpoint"] == "/guardrails/file"
    assert record["agent_key"] == "alice@co.com"  # header fallback
    assert record["metadata"]["device_id"] == "LAPTOP-42"
    assert record["metadata"]["file"]["name"] == "leak.txt"
    assert record["metadata"]["session_id"] == "sess-1"
    assert record["input_text"].startswith("file:leak.txt:")


def test_oversize_leaves_audit_trace(client, monkeypatch):
    """Oversize attachments are the most likely bulk-exfil shape — the 413
    must still write an audit entry before rejecting."""
    from storage import tenant_store
    import storage.audit_log as audit_log_mod

    r = FakeRedis()
    monkeypatch.setattr(tenant_store, "_get_redis", lambda: r)
    monkeypatch.setattr(tenant_store, "resolve_request_tenant_id",
                        lambda req: "bankco", raising=False)

    async def sync_log(self, entry):
        self._write_sync(entry)
    monkeypatch.setattr(audit_log_mod.AuditLogger, "log", sync_log)

    monkeypatch.setenv("SHIELD_FILE_MAX_BYTES", "100")
    resp = _upload(client, "dump.zip", b"x" * 200,
                   headers={"X-Device-Id": "LAPTOP-42"})
    assert resp.status_code == 413

    entries = r.zsets.get("audit:bankco", [])
    assert entries, "oversize rejection must be audited"
    record = json.loads(entries[-1][0])
    assert record["metadata"]["file"]["oversize"] is True
    assert record["metadata"]["file"]["name"] == "dump.zip"


if __name__ == "__main__":
    import sys
    sys.exit(pytest.main([__file__, "-v"]))
